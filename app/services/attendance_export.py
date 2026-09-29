"""Whole-term attendance workbook for one course (teacher export).

Columns are every date in the term whose weekday is in the course schedule;
rows are the enrolled students. Make-up sessions outside the schedule and
holidays are not modelled yet: a scheduled date with no session stays blank.
"""
import io
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TZ = ZoneInfo("Asia/Bangkok")
PAGE = 1000  # PostgREST returns at most this many rows per request
ON_TIME, LATE, ABSENT = "ตรงเวลา", "Late", "ไม่มา"
STATUS_LABELS = {"present": ON_TIME, "manual": ON_TIME, "late": LATE, "absent": ABSENT}
COLORS = {ON_TIME: ("C6EFCE", "006100"), LATE: ("FFEB9C", "9C5700"), ABSENT: ("FFC7CE", "9C0006")}
FONT = "Tahoma"  # covers Thai and Latin in one face


def read_all(make_query):
    """Every row of a query, page by page (one call silently stops at PAGE)."""
    rows = []
    while True:
        page = make_query().range(len(rows), len(rows) + PAGE - 1).execute().data or []
        rows.extend(page)
        if len(page) < PAGE:
            return rows


def term_dates(start_date, weeks, weekdays):
    """Every scheduled class date in the term (weekdays: 0=Mon, as in schedules)."""
    days = (start_date + timedelta(days=i) for i in range(weeks * 7))
    return [d for d in days if d.weekday() in weekdays]


def cell_status(day, today, session, row):
    if day > today or session is None:
        return ""  # not yet, or no class that day (holiday, server down)
    if row is not None:
        return STATUS_LABELS.get(row.get("status"), "")
    return "" if session.get("is_open") else ABSENT  # still open: not decided


def local_date(timestamp):
    return datetime.fromisoformat(timestamp.replace("Z", "+00:00")).astimezone(TZ).date()


def collect(db, course_id, term, today):
    """(dates, students) for the course: students = [{student_id, name, marks}]."""
    start = date.fromisoformat(str(term["start_date"]))
    weekdays = {s["day_of_week"] for s in read_all(lambda: db.table("schedules")
                .select("day_of_week").eq("course_id", course_id).order("id"))}
    dates = term_dates(start, int(term["weeks"]), weekdays)
    end = start + timedelta(weeks=int(term["weeks"]))
    sessions = read_all(lambda: db.table("sessions").select("id, start_time, is_open")
                        .eq("course_id", course_id)
                        .gte("start_time", f"{start.isoformat()}T00:00:00+07:00")
                        .lt("start_time", f"{end.isoformat()}T00:00:00+07:00").order("id"))
    by_date = {}
    for s in sessions:
        by_date.setdefault(local_date(s["start_time"]), s)
    session_ids = [s["id"] for s in by_date.values()]
    marks = {}
    if session_ids:
        for a in read_all(lambda: db.table("attendance").select("session_id, student_id, status")
                          .in_("session_id", session_ids).order("id")):
            marks[(a["session_id"], a["student_id"])] = a
    enrolled = read_all(lambda: db.table("course_enrollments")
                        .select("student_id, users(full_name, student_id)")
                        .eq("course_id", course_id).order("id"))
    students = []
    for e in enrolled:
        user = e.get("users") or {}
        row = [cell_status(d, today, by_date.get(d),
                           marks.get((by_date[d]["id"], e["student_id"])) if d in by_date else None)
               for d in dates]
        students.append({"student_id": user.get("student_id") or "",
                         "name": user.get("full_name") or "", "marks": row})
    students.sort(key=lambda s: (s["student_id"], s["name"]))
    return dates, students


def _text(cell, value):
    # openpyxl stores any string starting with "=" as a formula; names are text.
    cell.value = value
    cell.data_type = "s"
    return cell


def build_workbook(course, term, dates, students):
    wb = Workbook()
    ws = wb.active
    ws.title = (course.get("code") or "attendance")[:31]
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    last_col = max(2 + len(dates), 3)
    last_letter = get_column_letter(last_col)

    section = f" sec {course['section']}" if course.get("section") else ""
    ws.merge_cells(f"A1:{last_letter}1")
    _text(ws["A1"], f"ตารางเช็คชื่อ — {course.get('code', '')}{section} {course.get('name', '')}".strip())
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws.merge_cells(f"A2:{last_letter}2")
    start = date.fromisoformat(str(term["start_date"]))
    end = start + timedelta(weeks=int(term["weeks"]) - 1, days=6)
    thai = lambda d: f"{d.day}/{d.month}/{d.year}"  # same style as the date headers
    _text(ws["A2"], f"ภาคเรียน {term['name']} · {term['weeks']} สัปดาห์ ({thai(start)} – {thai(end)})")
    ws["A2"].font = Font(name=FONT, size=10, color="595959")

    header_row, first = 4, 5
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col, value in enumerate(["รหัสนักศึกษา", "ชื่อนักเรียน"] + dates, start=1):
        cell = ws.cell(row=header_row, column=col, value=value)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = center
        cell.border = box
        if isinstance(value, date):
            cell.number_format = "d/m/yyyy"

    for r, student in enumerate(students, start=first):
        _text(ws.cell(row=r, column=1), student["student_id"]).alignment = center
        _text(ws.cell(row=r, column=2), student["name"]).alignment = Alignment(vertical="center")
        for c, mark in enumerate(student["marks"], start=3):
            ws.cell(row=r, column=c, value=mark or None).alignment = center
        for c in range(1, 3 + len(dates)):
            ws.cell(row=r, column=c).border = box
            ws.cell(row=r, column=c).font = Font(name=FONT)

    last_row = first + max(len(students), 1) - 1
    if dates:
        status_range = f"C{first}:{get_column_letter(2 + len(dates))}{last_row}"
        dv = DataValidation(type="list", formula1='"' + ",".join(COLORS) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="สถานะไม่ถูกต้อง",
                            error="กรุณาเลือก ตรงเวลา, Late หรือ ไม่มา จากรายการ")
        dv.add(status_range)
        ws.add_data_validation(dv)
        for status, (fill, text) in COLORS.items():
            ws.conditional_formatting.add(status_range, CellIsRule(
                operator="equal", formula=[f'"{status}"'],
                fill=PatternFill("solid", fgColor=fill, bgColor=fill), font=Font(color=text, bold=True)))

    legend = last_row + 2
    ws.cell(row=legend, column=1, value="สถานะ").font = Font(name=FONT, bold=True)
    notes = {ON_TIME: "มาเรียน (รวมที่อาจารย์เช็คชื่อให้)", LATE: "มาสาย (เช็คชื่อหลังเริ่มคาบเกิน 15 นาที)",
             ABSENT: "ขาดเรียน / ไม่ได้เช็คชื่อในคาบที่ปิดแล้ว"}
    for i, (status, (fill, text)) in enumerate(COLORS.items()):
        cell = ws.cell(row=legend + i, column=2, value=status)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name=FONT, color=text, bold=True)
        cell.alignment = center
        cell.border = box
        ws.cell(row=legend + i, column=3, value=notes[status]).font = Font(name=FONT, color="595959")
    ws.cell(row=legend + 3, column=2,
            value="ช่องว่าง = ยังไม่ถึงวันเรียน หรือไม่มีคาบในวันนั้น · แก้ในไฟล์นี้ไม่เปลี่ยนข้อมูลในระบบ"
            ).font = Font(name=FONT, italic=True, color="7F7F7F")

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 26
    for c in range(3, 3 + len(dates)):
        ws.column_dimensions[get_column_letter(c)].width = 14  # fits "13/10/2026" in bold
    ws.row_dimensions[header_row].height = 22
    ws.freeze_panes = f"C{first}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
