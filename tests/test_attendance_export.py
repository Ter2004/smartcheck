import io
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook

import app
from app.routes import admin, auth, student, teacher
from app.services import attendance_export as export

TERM = {"name": "1/2569", "start_date": "2026-09-01", "weeks": 16}
COURSE = {"id": "c1", "code": "TEST01", "name": "Test class", "section": "001", "teacher_id": "t1"}


class Query:
    """Chainable stand-in for a PostgREST query over a list of rows."""

    def __init__(self, rows):
        self.rows, self.start, self.stop, self.single = list(rows), 0, None, False

    def __getattr__(self, name):  # select/eq/gte/lt/in_/order/limit: no-op filters
        return lambda *args, **kwargs: self

    def range(self, start, stop):
        self.start, self.stop = start, stop + 1
        return self

    def maybe_single(self):
        self.single = True
        return self

    def execute(self):
        rows = self.rows[self.start:self.stop]
        return SimpleNamespace(data=(rows[0] if rows else None) if self.single else rows)


class Database:
    def __init__(self, **tables):
        self.tables = tables

    def table(self, name):
        return Query(self.tables.get(name, []))


def enrolled(*people):
    return [{"student_id": uid, "users": {"student_id": sid, "full_name": name}}
            for uid, sid, name in people]


class TermAndStatusTests(unittest.TestCase):
    def test_sixteen_tuesdays_from_the_first_of_september(self):
        dates = export.term_dates(date(2026, 9, 1), 16, {1})
        self.assertEqual(len(dates), 16)
        self.assertEqual((dates[0], dates[-1]), (date(2026, 9, 1), date(2026, 12, 15)))

    def test_two_class_days_a_week_double_the_columns(self):
        self.assertEqual(len(export.term_dates(date(2026, 9, 1), 16, {1, 3})), 32)

    def test_cell_rules(self):
        today, past, future = date(2026, 9, 29), date(2026, 9, 22), date(2026, 10, 6)
        closed, open_ = {"is_open": False}, {"is_open": True}
        cases = [
            (future, closed, {"status": "present"}, ""),
            (past, None, None, ""),                       # no class that day
            (past, closed, {"status": "present"}, "ตรงเวลา"),
            (past, closed, {"status": "manual"}, "ตรงเวลา"),
            (past, closed, {"status": "late"}, "Late"),
            (past, closed, {"status": "absent"}, "ไม่มา"),
            (past, closed, None, "ไม่มา"),                # closed without check-in
            (today, open_, None, ""),                     # still open: undecided
        ]
        for day, session, row, expected in cases:
            with self.subTest(day=day, session=session, row=row):
                self.assertEqual(export.cell_status(day, today, session, row), expected)

    def test_read_all_follows_pages_past_the_server_limit(self):
        rows = [{"id": i} for i in range(2500)]
        self.assertEqual(len(export.read_all(lambda: Query(rows))), 2500)


class CollectAndWorkbookTests(unittest.TestCase):
    def db(self):
        return Database(
            schedules=[{"day_of_week": 1}],
            sessions=[{"id": "s1", "start_time": "2026-09-01T03:30:00+00:00", "is_open": False},
                      {"id": "s2", "start_time": "2026-09-08T03:30:00+00:00", "is_open": False}],
            attendance=[{"session_id": "s1", "student_id": "u1", "status": "late"},
                        {"session_id": "s2", "student_id": "u2", "status": "manual"}],
            course_enrollments=enrolled(("u2", "6630000002", "=HYPERLINK(\"http://x\")"),
                                        ("u1", "6630000001", "สมชาย ใจดี")),
        )

    def test_collect_maps_sessions_to_dates_and_sorts_students(self):
        dates, students = export.collect(self.db(), "c1", TERM, date(2026, 9, 16))
        self.assertEqual(len(dates), 16)
        first, second = students
        self.assertEqual(first["student_id"], "6630000001")
        # 1/9 late, 8/9 closed without check-in, 15/9 no session, later dates future.
        self.assertEqual(first["marks"][:4], ["Late", "ไม่มา", "", ""])
        self.assertEqual(second["marks"][:2], ["ไม่มา", "ตรงเวลา"])

    def test_workbook_layout_dropdown_colours_and_text_names(self):
        dates, students = export.collect(self.db(), "c1", TERM, date(2026, 9, 16))
        ws = load_workbook(io.BytesIO(export.build_workbook(COURSE, TERM, dates, students))).active
        self.assertIn("TEST01 sec 001", ws["A1"].value)
        self.assertEqual([ws.cell(4, c).value for c in (1, 2)], ["รหัสนักศึกษา", "ชื่อนักเรียน"])
        self.assertEqual(ws.cell(4, 3).value.date(), date(2026, 9, 1))
        self.assertEqual(ws.cell(4, 18).value.date(), date(2026, 12, 15))
        name = ws["B6"]
        self.assertEqual((name.value, name.data_type), ("=HYPERLINK(\"http://x\")", "s"))
        self.assertEqual(ws["C5"].value, "Late")
        (validation,) = ws.data_validations.dataValidation
        self.assertEqual(validation.formula1, '"ตรงเวลา,Late,ไม่มา"')
        self.assertEqual(str(validation.sqref), "C5:R6")
        rules = [r.formula for rng in ws.conditional_formatting for r in rng.rules]
        self.assertEqual(rules, [['"ตรงเวลา"'], ['"Late"'], ['"ไม่มา"']])


class ExportRouteTests(unittest.TestCase):
    def setUp(self):
        templates = Path(teacher.__file__).resolve().parents[1] / "templates"
        self.web = Flask(__name__, template_folder=str(templates))
        self.web.config.update(SECRET_KEY="test", RATELIMIT_ENABLED=False)
        app.limiter.init_app(self.web)
        self.web.register_blueprint(auth.auth_bp)
        self.web.register_blueprint(admin.admin_bp, url_prefix="/admin")
        self.web.register_blueprint(teacher.teacher_bp, url_prefix="/teacher")
        self.web.register_blueprint(student.student_bp, url_prefix="/student")
        self.client = self.web.test_client()
        with self.client.session_transaction() as session:
            session.update(user_id="t1", user_role="teacher")

    def use(self, **tables):
        patcher = patch.object(teacher, "supabase_admin", Database(**tables))
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_owner_downloads_an_xlsx(self):
        self.use(courses=[COURSE], terms=[TERM], schedules=[{"day_of_week": 1}],
                 course_enrollments=enrolled(("u1", "6630000001", "สมชาย ใจดี")))
        response = self.client.get("/teacher/course/c1/export")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.mimetype)
        self.assertIn("attendance_TEST01_sec001_1-2569.xlsx", response.headers["Content-Disposition"])
        load_workbook(io.BytesIO(response.data))

    def test_other_teachers_course_is_refused(self):
        self.use(courses=[dict(COURSE, teacher_id="t2")], terms=[TERM])
        response = self.client.get("/teacher/course/c1/export")
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.location.endswith("/teacher/dashboard"))

    def test_missing_term_explains_instead_of_exporting(self):
        self.use(courses=[COURSE], terms=[])
        response = self.client.get("/teacher/course/c1/export")
        self.assertEqual(response.status_code, 302)
        with self.client.session_transaction() as session:
            self.assertIn("ยังไม่ได้ตั้งวันเปิดเทอม", str(session.get("_flashes")))

    def test_per_session_csv_export_is_gone(self):
        self.assertEqual(self.client.get("/teacher/session/s1/export").status_code, 404)


if __name__ == "__main__":
    unittest.main()
